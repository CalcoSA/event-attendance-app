from collections.abc import Iterable, Mapping
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

HEADERS = (
    "Cédula", "Nombre", "Día programado", "Acompañantes programados", "Día real", "Bus",
    "Titular asistió", "Acompañantes reales", "Total asistentes", "Cédula usuario que registró",
    "Nombre usuario que registró", "Fecha y hora",
)
WIDTHS = (22, 38, 29, 25, 29, 10, 17, 23, 20, 29, 38, 25)


def build_workbook(rows: Iterable[Mapping], buses: Iterable[Mapping]) -> BytesIO:
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets_by_bus_id = {}
    for bus in buses:
        sheet = workbook.create_sheet(bus["display_name"])
        sheets_by_bus_id[bus["bus_id"]] = sheet
        sheet.append(HEADERS)
        sheet.freeze_panes = "A2"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.sheet_properties.tabColor = "52362A"
        sheet.row_dimensions[1].height = 34
        for column, width in enumerate(WIDTHS, 1):
            cell = sheet.cell(1, column)
            cell.font = Font(name="Calibri", bold=True, color="FFF9E9")
            cell.fill = PatternFill("solid", fgColor="52362A")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            sheet.column_dimensions[get_column_letter(column)].width = width
    for row in rows:
        number = int(row["bus_number"])
        if not 1 <= number <= 16:
            raise ValueError("La asistencia contiene un bus fuera del rango permitido.")
        sheet = sheets_by_bus_id[row["bus_id"]]
        moment = row["checked_in_at"]
        if isinstance(moment, datetime) and moment.tzinfo is not None:
            from .services import BOGOTA
            moment = moment.astimezone(BOGOTA).replace(tzinfo=None)
        values = (
            str(row["document"]), row["full_name"], row["planned_day"], int(row["planned_companion_count"]),
            row["actual_day"], row["display_name"], "Sí" if row["titular_present"] else "No",
            int(row["actual_companions"]), int(row["total_present"]), str(row["registered_by_document"]),
            row["registered_by_name"], moment,
        )
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            cell.alignment = Alignment(vertical="center")
            if isinstance(cell.value, str):
                # Explicit text cells prevent formulas even for =,+,-,@ input.
                cell.value = ILLEGAL_CHARACTERS_RE.sub("", cell.value)
                cell.data_type = "s"
            if sheet.max_row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="FFF9E9")
        sheet.cell(sheet.max_row, 1).number_format = "@"
        sheet.cell(sheet.max_row, 10).number_format = "@"
        sheet.cell(sheet.max_row, 12).number_format = "yyyy-mm-dd hh:mm:ss"
    for sheet in workbook:
        sheet.auto_filter.ref = f"A1:L{sheet.max_row}"
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream
