from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from types import SimpleNamespace
import unittest
from unittest.mock import patch
from zoneinfo import ZoneInfo

from argon2 import PasswordHasher, Type
from fastapi.testclient import TestClient
import jwt
from openpyxl import load_workbook
from sqlalchemy.exc import OperationalError

from app import main
from app.database import get_db
from app.schemas import AttendanceSummary
from app.security import COOKIE_NAME, JWT_AUDIENCE, JWT_ISSUER, create_access_token
from app.services import AttendanceAlreadyExists
from tests.helpers import (
    attendance_body, database_session, previous_attendance, sample_buses, sample_event, sample_operator,
    test_settings,
)


class ApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.password = "isolated-test-password"
        cls.password_hash = PasswordHasher(type=Type.ID).hash(cls.password)

    def setUp(self):
        self.settings = test_settings()
        self.db = database_session()
        self.operator = sample_operator(password_hash=self.password_hash)
        self.db.get.return_value = self.operator
        self.app = main.create_app(self.settings)
        self.app.dependency_overrides[get_db] = lambda: self.db
        self.client = TestClient(self.app)
        self.client.__enter__()
        self.addCleanup(self.client.__exit__, None, None, None)
        self.post_headers = {"X-Requested-With": "XMLHttpRequest"}

    def authenticate(self):
        self.client.cookies.set(COOKIE_NAME, create_access_token(self.operator.user_id, self.settings))

    def test_login_uses_httponly_cookie_updates_last_login_and_logout_deletes_cookie(self):
        self.db.scalar.return_value = self.operator
        response = self.client.post(
            "/api/auth/login", headers=self.post_headers,
            json={"username": "0000000002", "password": self.password},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["document"], "0000000002")
        self.assertEqual(set(response.json()), {"user_id", "document", "full_name", "role"})
        cookie = response.headers["set-cookie"].lower()
        self.assertIn("httponly", cookie)
        self.assertIn("samesite=lax", cookie)
        self.assertIn("path=/", cookie)
        self.assertNotIn("; secure", cookie)
        self.assertIsNotNone(self.operator.last_login_at)
        self.assertIsNone(self.operator.last_login_at.tzinfo)
        self.db.commit.assert_called_once()
        self.assertEqual(self.client.get("/api/auth/me").status_code, 200)
        logout = self.client.post("/api/auth/logout", headers=self.post_headers)
        self.assertEqual(logout.status_code, 200)
        self.assertIn("Max-Age=0", logout.headers["set-cookie"])
        self.assertEqual(self.client.get("/api/auth/me").status_code, 401)

    def test_login_failures_have_one_generic_message(self):
        for operator in (
            None,
            sample_operator(password_hash=self.password_hash, is_active=False),
            sample_operator(password_hash="invalid-argon2id-hash"),
        ):
            with self.subTest(operator_exists=operator is not None):
                self.db.scalar.return_value = operator
                response = self.client.post(
                    "/api/auth/login", headers=self.post_headers,
                    json={"username": "0000000002", "password": self.password},
                )
                self.assertEqual(response.status_code, 401)
                self.assertEqual(response.json(), {"detail": "Usuario o contraseña incorrectos."})
                self.assertNotIn("set-cookie", response.headers)
        self.db.commit.assert_not_called()

    def test_sensitive_endpoints_require_authentication(self):
        for url in (
            "/api/auth/me", "/api/event/context", "/api/participants",
            "/api/participants/0000000001", "/api/export/attendance.xlsx",
        ):
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 401)
        self.assertEqual(self.client.post(
            "/api/attendance", headers=self.post_headers, json=attendance_body()
        ).status_code, 401)
        self.db.execute.assert_not_called()
        self.db.scalar.assert_not_called()

    def test_expired_token_and_inactive_user_cannot_access_data(self):
        expired = datetime.now(timezone.utc) - timedelta(minutes=1)
        token = jwt.encode(
            {"sub": "901", "iat": expired - timedelta(minutes=1), "exp": expired,
             "iss": JWT_ISSUER, "aud": JWT_AUDIENCE},
            self.settings.JWT_SECRET.get_secret_value(), algorithm="HS256",
        )
        self.client.cookies.set(COOKIE_NAME, token)
        self.assertEqual(self.client.get("/api/auth/me").status_code, 401)
        self.db.get.assert_not_called()
        self.authenticate()
        self.operator.is_active = False
        self.assertEqual(self.client.get("/api/auth/me").status_code, 403)

    def test_write_requests_need_csrf_header_and_validation_never_echoes_password(self):
        response = self.client.post(
            "/api/auth/login", json={"username": "0000000002", "password": self.password}
        )
        self.assertEqual(response.status_code, 403)
        self.db.scalar.assert_not_called()
        secret_input = "sensitive-test-input-" * 20
        response = self.client.post(
            "/api/auth/login", headers=self.post_headers,
            json={"username": "0000000002", "password": secret_input},
        )
        self.assertEqual(response.status_code, 422)
        self.assertNotIn(secret_input, response.text)
        self.assertTrue(all("input" not in error for error in response.json()["detail"]))

    def test_attendance_uses_authenticated_operator_returns_201_and_duplicate_409(self):
        self.authenticate()
        summary = AttendanceSummary.model_validate(previous_attendance(
            bus_id=516, bus_number=16, display_name="No aplica"
        ))
        with patch.object(main, "get_event", return_value=sample_event()), patch.object(
            main, "create_attendance", return_value=summary
        ) as create:
            response = self.client.post("/api/attendance", headers=self.post_headers, json=attendance_body(bus_id=516))
            self.assertEqual(response.status_code, 201)
            self.assertEqual(response.json()["bus_id"], 516)
            self.assertEqual(response.json()["bus_number"], 16)
            self.assertEqual(response.json()["display_name"], "No aplica")
            self.assertIs(create.call_args.args[2], self.operator)
            self.assertEqual(create.call_args.args[3].participant_id, 701)
            self.assertEqual(create.call_args.args[3].bus_id, 516)
            create.side_effect = AttendanceAlreadyExists(summary)
            duplicate = self.client.post("/api/attendance", headers=self.post_headers, json=attendance_body(bus_id=516))
        self.assertEqual(duplicate.status_code, 409)
        self.assertEqual(duplicate.json()["detail"], "Esta persona ya fue registrada previamente.")
        self.assertEqual(duplicate.json()["attendance"]["total_present"], 2)
        self.assertEqual(duplicate.json()["attendance"]["display_name"], "No aplica")
        self.assertEqual(duplicate.json()["attendance"]["bus_id"], 516)

    def test_context_reads_real_identifiers_and_suggests_bogota_day(self):
        self.authenticate()
        self.db.scalars.side_effect = [
            [SimpleNamespace(event_day_id=601, event_date=date(2026, 10, 6), day_label="Martes 6 de octubre"),
             SimpleNamespace(event_day_id=602, event_date=date(2026, 10, 7), day_label="Miércoles 7 de octubre")],
            [SimpleNamespace(**sample_buses()[0]), SimpleNamespace(**sample_buses()[-1])],
        ]
        with patch.object(main, "get_event", return_value=sample_event()), patch.object(
            main, "now_local", return_value=datetime(2026, 10, 7, 9, tzinfo=ZoneInfo("America/Bogota"))
        ):
            response = self.client.get("/api/event/context")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["suggested_day_id"], 602)
        self.assertEqual(response.json()["buses"], [
            {"bus_id": 501, "bus_number": 1, "display_name": "Bus 1"},
            {"bus_id": 516, "bus_number": 16, "display_name": "No aplica"},
        ])

    def test_export_stream_is_a_real_workbook_and_is_never_cached(self):
        self.authenticate()
        self.db.execute.return_value.mappings.return_value.all.return_value = sample_buses()
        with patch.object(main, "get_event", return_value=sample_event()), patch.object(
            main, "export_rows", return_value=[]
        ), patch.object(main, "now_local", return_value=datetime(2026, 10, 7, 9, 30)):
            response = self.client.get("/api/export/attendance.xlsx")
        self.assertEqual(response.status_code, 200)
        self.assertIn("asistencia_fiesta_ninos_2026_20261007_0930.xlsx", response.headers["content-disposition"])
        self.assertEqual(response.headers["cache-control"], "no-store")
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(len(workbook.worksheets), 16)
        self.assertEqual(workbook.sheetnames[-1], "No aplica")
        self.assertNotIn("Bus 16", workbook.sheetnames)

    def test_database_failures_return_controlled_error_without_connection_details(self):
        secret = "do-not-expose-test-connection-password"
        self.db.scalar.side_effect = OperationalError("SELECT private", {"password": secret}, Exception(secret))
        response = self.client.post(
            "/api/auth/login", headers=self.post_headers,
            json={"username": "0000000002", "password": self.password},
        )
        self.assertEqual(response.status_code, 500)
        self.assertNotIn(secret, response.text)
        self.assertEqual(response.headers["cache-control"], "no-store")

    def test_secure_cookie_and_cors_origins_are_explicit(self):
        settings = test_settings(COOKIE_SECURE=True, CORS_ORIGINS=["https://allowed.example.invalid"])
        application = main.create_app(settings)
        application.dependency_overrides[get_db] = lambda: self.db
        self.db.scalar.return_value = self.operator
        with TestClient(application, base_url="https://testserver") as client:
            response = client.post(
                "/api/auth/login", headers=self.post_headers,
                json={"username": "0000000002", "password": self.password},
            )
            self.assertIn("; Secure", response.headers["set-cookie"])
            for origin, expected in (
                ("https://allowed.example.invalid", 200),
                ("https://rejected.example.invalid", 400),
            ):
                preflight = client.options("/api/attendance", headers={
                    "Origin": origin, "Access-Control-Request-Method": "POST",
                    "Access-Control-Request-Headers": "X-Requested-With, Content-Type",
                })
                self.assertEqual(preflight.status_code, expected)
                self.assertNotEqual(preflight.headers.get("access-control-allow-origin"), "*")


if __name__ == "__main__":
    unittest.main()
