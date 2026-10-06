from datetime import date, datetime
import unittest

from openpyxl import load_workbook

from app.exports import build_workbook
from app import services
from app.models import Bus
from tests.helpers import database_session, sample_buses


class ExportTests(unittest.TestCase):
    def test_empty_export_has_exactly_sixteen_usable_bus_sheets(self):
        workbook = load_workbook(build_workbook([], sample_buses()))
        self.assertEqual(workbook.sheetnames, [f"Bus {number}" for number in range(1, 16)] + ["No aplica"])
        for sheet in workbook:
            with self.subTest(sheet=sheet.title):
                self.assertEqual(sheet.max_row, 1)
                self.assertEqual(sheet.max_column, 12)
                self.assertEqual(sheet.freeze_panes, "A2")
                self.assertEqual(sheet.auto_filter.ref, "A1:L1")
                self.assertTrue(all(cell.value for cell in sheet[1]))

    def test_data_is_grouped_by_bus_and_documents_remain_text(self):
        row = {
            "document": "0000000001",
            "full_name": "PERSONA DE PRUEBA",
            "planned_date": date(2026, 10, 6),
            "planned_day": "Martes 6 de octubre",
            "planned_companion_count": 2,
            "actual_date": date(2026, 10, 7),
            "actual_day": "Miércoles 7 de octubre",
            "bus_id": 516,
            "bus_number": 16,
            "display_name": "No aplica",
            "titular_present": 1,
            "actual_companions": 3,
            "total_present": 4,
            "registered_by_document": "0000000002",
            "registered_by_name": "OPERADOR DE PRUEBA",
            "checked_in_at": datetime(2026, 10, 7, 9, 12, 34),
        }
        workbook = load_workbook(build_workbook([row], sample_buses()))
        self.assertEqual(workbook["Bus 1"].max_row, 1)
        sheet = workbook["No aplica"]
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet["A2"].value, "0000000001")
        self.assertEqual(sheet["A2"].data_type, "s")
        self.assertEqual(sheet["B2"].value, "PERSONA DE PRUEBA")
        self.assertEqual(sheet["D2"].value, 2)
        self.assertEqual(sheet["F2"].value, "No aplica")
        self.assertEqual(sheet["F2"].data_type, "s")
        self.assertEqual(sheet["H2"].value, 3)
        self.assertEqual(sheet["I2"].value, 4)
        self.assertEqual(sheet["J2"].value, "0000000002")
        self.assertEqual(sheet["J2"].data_type, "s")
        self.assertEqual(sheet.auto_filter.ref, "A1:L2")

    def test_spreadsheet_formula_like_names_cannot_be_executed(self):
        row = {
            "document": "0000000001",
            "full_name": '=HYPERLINK("https://example.invalid", "prueba")',
            "planned_day": "Martes 6 de octubre",
            "planned_companion_count": 0,
            "actual_day": "Martes 6 de octubre",
            "bus_id": 501,
            "bus_number": 1,
            "display_name": "Bus 1",
            "titular_present": 1,
            "actual_companions": 0,
            "total_present": 1,
            "registered_by_document": "0000000002",
            "registered_by_name": "=1+1",
            "checked_in_at": datetime(2026, 10, 6, 9),
        }
        workbook = load_workbook(build_workbook([row], sample_buses()), data_only=False)
        self.assertEqual(workbook["Bus 1"]["F2"].value, "Bus 1")
        for column in ("B", "K"):
            self.assertNotEqual(workbook["Bus 1"][f"{column}2"].data_type, "f")

    def test_sheet_names_follow_catalog_display_names(self):
        buses = sample_buses()
        buses[-1]["display_name"] = "Sin transporte"
        workbook = load_workbook(build_workbook([], buses))
        self.assertEqual(workbook.sheetnames[-1], "Sin transporte")
        self.assertNotIn("Bus 16", workbook.sheetnames)

    def test_export_query_selects_bus_identity_number_and_display_name(self):
        db = database_session()
        rows = [{"bus_id": 516, "bus_number": 16, "display_name": "No aplica"}]
        db.execute.return_value.mappings.return_value.all.return_value = rows

        self.assertEqual(services.export_rows(db, 801), rows)

        statement = db.execute.call_args.args[0]
        for column in (Bus.bus_id, Bus.bus_number, Bus.display_name):
            with self.subTest(column=column.key):
                self.assertTrue(statement.selected_columns[column.key].shares_lineage(column))


if __name__ == "__main__":
    unittest.main()
